"""
=============================================================================
Foxe Studio — Parser Neraca Keuangan (File Neraca Google Drive)
=============================================================================
Parser khusus untuk mengurai pos pengeluaran dari sheet 'September 2026'
pada section 'Detail' (Kolom P s.d. U) sesuai arahan pemilik studio:
- Membaca baris per tanggal menurun ke bawah.
- Mengklasifikasikan otomatis ke:
  * COGS (Biaya Langsung Produksi / Sesi Foto)
  * OPEX (Biaya Operasional Studio)
  * NON-P&L / Lainnya
=============================================================================
"""

import os
import openpyxl

NERACA_SHEET_ID = "1dvnCNyfZI5z-12081XJjGCVLtMaQpUStYT61qU3orKM"

def classify_expense(desc):
    d = str(desc or "").lower().strip()
    
    # 1. COGS (Biaya Langsung Produksi / Paket / Sesi Foto)
    if any(k in d for k in ["cetak", "print"]):
        return "COGS", "Cetak Foto & Photo Paper"
    if "frame" in d or "bingkai" in d or "album" in d:
        return "COGS", "Frame / Album / Packaging"
    if any(k in d for k in ["batrei", "baterai", "rent batrei", "sewa batrei"]):
        return "COGS", "Properti / Consumable Sesi"
    if "rent equipment" in d or "sewa alat" in d:
        return "COGS", "Outsource / Freelancer Produksi"
    if any(k in d for k in ["freelance", "ruslan", "camera", "fg "]):
        return "COGS", "Outsource / Freelancer Produksi"
        
    # 2. OPEX (Biaya Operasional Studio Rutin)
    if "listrik" in d:
        return "OPEX", "Listrik & Air"
    if "wifi" in d or "internet" in d:
        return "OPEX", "Internet & Telekomunikasi"
    if "ads" in d or "iklan" in d:
        return "OPEX", "Marketing / Ads / KOL"
    if any(k in d for k in ["makan", "sarapan", "konsumsi", "snack"]):
        return "OPEX", "Konsumsi Crew"
    if "adobe" in d or "software" in d or "subscription" in d:
        return "OPEX", "Software / Subscription"
    if any(k in d for k in ["cat limbo", "paralon", "maintenance", "dop", "perbaikan"]):
        return "OPEX", "Maintenance Studio & Equipment"
    if any(k in d for k in ["galon", "parfum", "cleaning", "tissue", "kebersihan"]):
        return "OPEX", "Office & Cleaning Supplies"
    if "gaji" in d:
        return "OPEX", "Gaji Admin & Fotografer"
        
    # 3. Non-PL / Lainnya
    if "bon" in d:
        return "NON-P&L", "Prive / Owner Draw"
    if "pajak" in d:
        return "OPEX", "Pajak & Perizinan"
    if "amal" in d:
        return "Biaya Lainnya", "Biaya Lainnya"
        
    return "OPEX", "OPEX Lainnya"

def clean_num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("Rp", "").strip()
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except:
        return 0.0

def parse_neraca(filepath="file_neraca.xlsx", sheet_name="September 2026"):
    if not os.path.exists(filepath):
        print(f"[WARN] File neraca {filepath} tidak ditemukan.")
        return []

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        # Fallback to first sheet
        sheet_name = wb.sheetnames[0]
        
    ws = wb[sheet_name]

    current_tgl = None
    expenses = []
    idx = 1

    # Section Detail terletak pada kolom:
    # Col 16 (P): Tgl
    # Col 17 (Q): Saldo Masuk
    # Col 18 (R): Saldo Keluar
    # Col 19 (S): Balance
    # Col 20 (T): Keterangan Masuk
    # Col 21 (U): Keterangan Keluar

    for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if r_idx < 10 or len(row) < 21:
            continue
            
        tgl_val = row[15]
        keluar_val = row[17]
        ket_keluar = row[20]
        
        # Tanggal merge/vertikal
        if tgl_val is not None and str(tgl_val).strip() != "":
            try:
                parsed_day = int(float(str(tgl_val).strip()))
                if 1 <= parsed_day <= 31:
                    current_tgl = parsed_day
            except:
                pass
                
        nominal = clean_num(keluar_val)
        desc = str(ket_keluar or "").strip()
        
        # Abaikan subtotal / angka numerik yang bukan keterangan
        if desc.replace(".", "").replace(",", "").isdigit() or desc.lower() in ["total", "nominal", "balance"]:
            continue
            
        if nominal > 0 and desc and current_tgl is not None:
            jenis, kategori = classify_expense(desc)
            tgl_iso = f"2026-09-{current_tgl:02d}"
            expenses.append({
                "id": f"ex_nrc_{idx}",
                "tanggal": tgl_iso,
                "bulan": "2026-09",
                "jenis": jenis,
                "kategori": kategori,
                "deskripsi": desc,
                "nilai": nominal
            })
            idx += 1

    print(f"[OK] File Neraca berhasil diurai: {len(expenses)} pos pengeluaran dari section Detail.")
    return expenses

if __name__ == "__main__":
    exps = parse_neraca()
    cogs_tot = sum(e["nilai"] for e in exps if e["jenis"] == "COGS")
    opex_tot = sum(e["nilai"] for e in exps if e["jenis"] == "OPEX")
    print(f"Total Pos: {len(exps)}")
    print(f"Total COGS: Rp {cogs_tot:,.0f}")
    print(f"Total OPEX: Rp {opex_tot:,.0f}")
    print(f"Total Biaya: Rp {(cogs_tot+opex_tot):,.0f}")

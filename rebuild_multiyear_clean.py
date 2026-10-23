import sys
sys.path.append('.')
import json
import os
import openpyxl
from build_multiyear_cache import classify_expense, clean_num, find_sheet_for_month, BULAN_NAMES

print("Loading file_neraca.xlsx...")
wb_formula = openpyxl.load_workbook('file_neraca.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('file_neraca.xlsx', data_only=True)
sheetnames = wb_val.sheetnames

with open('multiyear_cache.json', 'r', encoding='utf-8') as f:
    my = json.load(f)

# Re-parse 2022, 2023, 2024, 2025
for y in [2022, 2023, 2024, 2025]:
    year_str = str(y)
    print(f"\nProcessing Year {year_str}...")
    
    y_omz = 0
    y_cash = 0
    y_trf = 0
    y_cg = 0
    y_op = 0
    y_out = 0
    y_nt = 0
    
    new_months = []
    
    for m in range(1, 13):
        sname = find_sheet_for_month(sheetnames, year_str, m)
        if not sname or sname not in wb_val.sheetnames:
            print(f"  M{m:02d}: Not found")
            continue
            
        ws_f = wb_formula[sname]
        ws_v = wb_val[sname]
        
        # 1. Parse daily cash/transfer (rows 11..41)
        daily = []
        tot_c_daily = 0.0
        tot_t_daily = 0.0
        tot_in_daily = 0.0
        
        for r in range(11, 42):
            d_val = ws_v.cell(r, 2).value
            if d_val is None:
                continue
            try:
                day_int = int(float(str(d_val).strip()))
                if not (1 <= day_int <= 31):
                    continue
            except:
                continue
                
            c_val = clean_num(ws_v.cell(r, 3).value)
            q_val = clean_num(ws_v.cell(r, 4).value)
            t_val = clean_num(ws_v.cell(r, 5).value)
            f_val = clean_num(ws_v.cell(r, 6).value)
            if f_val == 0.0 and (c_val + q_val + t_val) > 0:
                f_val = c_val + q_val + t_val
                
            tot_c_daily += c_val
            tot_t_daily += (t_val + q_val)
            tot_in_daily += f_val
            
            daily.append({
                "tgl": f"{y:04d}-{m:02d}-{day_int:02d}",
                "day": day_int,
                "cash": round(c_val),
                "transfer": round(t_val + q_val),
                "total": round(f_val)
            })
            
        # Check sheet total row around row 42..45
        tot_c_sheet = 0.0
        tot_t_sheet = 0.0
        tot_in_sheet = 0.0
        for r in range(41, 46):
            b_val = str(ws_v.cell(r, 2).value or "").lower()
            if "total" in b_val or "jumlah" in b_val or ws_v.cell(r, 6).value is not None:
                c = clean_num(ws_v.cell(r, 3).value)
                q = clean_num(ws_v.cell(r, 4).value)
                t = clean_num(ws_v.cell(r, 5).value)
                f = clean_num(ws_v.cell(r, 6).value)
                if f > 0 or (c + t + q) > 0:
                    tot_c_sheet = c
                    tot_t_sheet = t + q
                    tot_in_sheet = f if f > 0 else (c + t + q)
                    break
                    
        omzet = tot_in_sheet if tot_in_sheet > 0 else tot_in_daily
        cash = tot_c_sheet if tot_c_sheet > 0 else tot_c_daily
        transfer = tot_t_sheet if tot_t_sheet > 0 else tot_t_daily
        if omzet == 0 and (cash + transfer) > 0:
            omzet = cash + transfer
            
        # 2. Parse expenses from row 11 up to the total row (=SUM)
        expenses = []
        tot_cogs = 0.0
        tot_opex = 0.0
        current_day = 1
        
        for r in range(11, 150):
            # Stop condition: total formula in col R
            f_val = str(ws_f.cell(r, 18).value or '')
            if f_val.startswith('=SUM') and (':' in f_val or '+' in f_val):
                break
                
            ket_val = str(ws_v.cell(r, 21).value or ws_v.cell(r, 20).value or ws_v.cell(r, 17).value or '').lower()
            if any(w in ket_val for w in ['bon manajemen', 'alas', 'folx', 'tambahan / balik', 'refund out']):
                continue
                
            p_val = ws_v.cell(r, 16).value
            if p_val is not None:
                try:
                    d = int(float(str(p_val).strip()))
                    if 1 <= d <= 31:
                        current_day = d
                except:
                    pass
                    
            out_val = clean_num(ws_v.cell(r, 18).value)
            desc_val = str(ws_v.cell(r, 21).value or ws_v.cell(r, 20).value or '').strip()
            
            if out_val == 0.0 and not desc_val:
                continue
            if desc_val.lower() in ['saldo keluar', 'keterangan', 'detail', 'nominal']:
                continue
                
            jenis, kat = classify_expense(desc_val)
            if jenis == 'COGS':
                tot_cogs += out_val
            else:
                tot_opex += out_val
                
            expenses.append({
                "tgl": f"{y:04d}-{m:02d}-{current_day:02d}",
                "day": current_day,
                "desc": desc_val or "Pengeluaran Operasional",
                "jenis": jenis,
                "kategori": kat,
                "nilai": round(out_val)
            })
            
        outcome = tot_cogs + tot_opex
        nett_profit = omzet - outcome
        margin = (nett_profit / omzet * 100.0) if omzet > 0 else 0.0
        ads = sum(e["nilai"] for e in expenses if e.get("kategori") == "Marketing / Ads / KOL" or "ads" in e.get("desc", "").lower() or "iklan" in e.get("desc", "").lower())
        
        m_obj = {
            "month": m,
            "name": BULAN_NAMES[m - 1],
            "omzet": round(omzet),
            "cash": round(cash),
            "transfer": round(transfer),
            "cogs": round(tot_cogs),
            "opex": round(tot_opex),
            "outcome": round(outcome),
            "nettProfit": round(nett_profit),
            "margin": round(margin, 2),
            "ads": round(ads),
            "daily": daily,
            "expenses": expenses,
            "sheetName": sname
        }
        new_months.append(m_obj)
        
        y_omz += omzet
        y_cash += cash
        y_trf += transfer
        y_cg += tot_cogs
        y_op += tot_opex
        y_out += outcome
        y_nt += nett_profit
        
        print(f"  M{m:02d} ({BULAN_NAMES[m-1]:10s}): Omzet={omzet:11,.0f} | COGS={tot_cogs:10,.0f} | OPEX={tot_opex:10,.0f} | Outflow={outcome:11,.0f} | ExpCount={len(expenses)}")

    y_margin = (y_nt / y_omz * 100.0) if y_omz > 0 else 0.0
    my["years"][year_str] = {
        "year": year_str,
        "omzet": round(y_omz),
        "cash": round(y_cash),
        "transfer": round(y_trf),
        "cogs": round(y_cg),
        "opex": round(y_op),
        "outcome": round(y_out),
        "nettProfit": round(y_nt),
        "margin": round(y_margin, 2),
        "months": new_months
    }
    print(f"Year {year_str} Total: Omzet={y_omz:,.0f} | COGS={y_cg:,.0f} | OPEX={y_op:,.0f} | Outflow={y_out:,.0f} | Nett={y_nt:,.0f} ({y_margin:.1f}%)")

# Save to multiyear_cache.json
with open("multiyear_cache.json", "w", encoding="utf-8") as f:
    json.dump(my, f, ensure_ascii=False, indent=2)
print("\nSaved multiyear_cache.json successfully.")

# Sync to foxe_full_state.json
with open("foxe_full_state.json", "r", encoding="utf-8") as f:
    fs = json.load(f)
fs["multiyearArsip"] = my
with open("foxe_full_state.json", "w", encoding="utf-8") as f:
    json.dump(fs, f, ensure_ascii=False, indent=2)
print("Synced to foxe_full_state.json successfully.")

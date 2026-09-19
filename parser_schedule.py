"""
=============================================================================
Foxe Studio — Parser Schedule Booking (File 2)
=============================================================================
Parser khusus untuk mengurai jadwal booking studio dari 'file2.xlsx' (Schedule).
Membaca sheet 'Tanggal 1' s.d. 'Tanggal 31':
1. Memindai 3 Blok Studio:
   - Studio 1: Kolom 1 s.d. 6  (No, Waktu, Nama, Paket, Admin, No.HP)
   - Studio 2: Kolom 7 s.d. 12 (No, Waktu, Nama, Paket, Admin, No.HP)
   - Studio 3: Kolom 13 s.d. 18 (No, Waktu, Nama, Paket, Admin, No.HP)
2. Normalisasi nama paket & penentuan harga:
   - Large Group dinamis: 'LG <n>' / 'Large Group <n>' = n × Rp25.000 (min 7)
   - Paket standar (Graduation, Couple, Photofox, Pas Foto, dsb.)
3. Mempertahankan booking berstatus manual: true jika terjadi bentrok.
=============================================================================
"""

import re
import datetime
import openpyxl

PRICING_TABLE = {
    "Graduation": 350000,
    "Graduation Premium": 500000,
    "Large Group": 25000, # per orang (min 7)
    "Single": 100000,
    "Couple A": 150000,
    "Couple B": 200000,
    "Couple C": 400000,
    "Photofox": 200000,
    "Photo Fox": 200000,
    "Family A": 350000,
    "Family B": 450000,
    "Maternity": 350000,
    "Studio Rent": 450000,
    "Pas Foto": 50000,
    "All File": 100000,
    "All Soft File": 50000,
    "Add Cetak": 25000,
    "Add Edit": 50000,
    "Packing": 20000
}

def str_time(v):
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime.datetime):
        return v.strftime("%H:%M")
    return str(v or "").strip()

def npak(p):
    s = str(p or "").strip()
    if s.lower() == "photo fox":
        return "Photofox"
    if s.lower().startswith("pass photo"):
        return "Pas Foto"
    return s

def parse_schedule(filepath="file2.xlsx", bulan="2026-09", existing_bookings=None):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    manual_map = {}
    if existing_bookings:
        for b in existing_bookings:
            if b.get("manual"):
                k = f"{b.get('tgl')}|{b.get('nama','').strip().lower()}|{b.get('waktu')}"
                manual_map[k] = b

    bookings = []
    b_idx = 1
    
    for day in range(1, 32):
        sname = f"Tanggal {day}"
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        tgl_str = f"{bulan}-{day:02d}"
        
        studios = [(1, "Studio 1"), (7, "Studio 2"), (13, "Studio 3")]
        for c_start, studio_name in studios:
            for r in range(3, ws.max_row + 1):
                nama_val = ws.cell(r, c_start + 2).value
                if not nama_val:
                    continue
                nama = str(nama_val).strip()
                if not nama or nama.lower() in ["nama", "none", "-"]:
                    continue
                
                waktu = str_time(ws.cell(r, c_start + 1).value)
                paket_raw = str(ws.cell(r, c_start + 3).value or "").strip()
                admin = str(ws.cell(r, c_start + 4).value or "").strip().upper()
                hp = str(ws.cell(r, c_start + 5).value or "").strip()
                
                # Cek manual override
                m_key = f"{tgl_str}|{nama.lower()}|{waktu}"
                if m_key in manual_map:
                    bookings.append(manual_map[m_key])
                    continue
                
                # Hitung harga paket
                paket_std = npak(paket_raw)
                harga = 0.0
                jml_orang = 1
                
                # Cek Large Group (misal: "LG 10", "Large Group 15")
                lg_match = re.search(r'(?:lg|large\s*group)\s*(\d+)', paket_raw, re.IGNORECASE)
                if lg_match:
                    paket_std = "Large Group"
                    jml_orang = int(lg_match.group(1))
                    harga = max(7, jml_orang) * 25000.0
                elif paket_std in PRICING_TABLE:
                    harga = float(PRICING_TABLE[paket_std])
                elif "wisuda" in paket_raw.lower() or "grad" in paket_raw.lower():
                    paket_std = "Graduation"
                    harga = 350000.0
                elif "couple" in paket_raw.lower():
                    paket_std = "Couple A"
                    harga = 150000.0
                elif "group" in paket_raw.lower():
                    paket_std = "Large Group"
                    harga = 175000.0
                else:
                    harga = 0.0
                
                bookings.append({
                    "id": f"sc_bk_{b_idx}",
                    "tgl": tgl_str,
                    "waktu": waktu,
                    "studio": studio_name,
                    "studioNama": studio_name,
                    "nama": nama,
                    "paket": paket_std,
                    "paketRaw": paket_raw,
                    "jmlOrang": jml_orang,
                    "harga": harga,
                    "admin": admin,
                    "noHp": hp,
                    "manual": False
                })
                b_idx += 1

    wb.close()
    return bookings

if __name__ == "__main__":
    res = parse_schedule()
    print(f"Hasil Parser Schedule: {len(res)} sesi foto terdata.")
